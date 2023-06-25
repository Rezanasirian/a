import json
import pandas as pd
import numpy as np

from openpyxl import Workbook
import docx
from docx.shared import Pt, Mm
from docx.enum.style import WD_STYLE_TYPE


with open('C:/Users/Lenovo/Desktop/dundas7.txt',encoding='utf8') as f:
    file_contents = f.read()
    parsed_json = json.loads(file_contents)
    adaptor = parsed_json['Adapters']
    print(parsed_json.keys())
class Chart:
    def __init__(self):
        self.metricsets=[]
        self.measurs = []
        self.action = []
        self.name = []
        self.type = []
        self.script= []

    def readmetricset(self):
        if 'MetricSetBindings' in abj.keys():
            for metricset_list in abj['MetricSetBindings']:
                self.metricsets.append(metricset_list['FriendlyName'])

    def readmeasures(self):
        if 'MetricSetBindings' in abj.keys():
            for metricset_list in abj['MetricSetBindings']:
                for metricset in metricset_list['Bindings']:
                    if 'Measures'in  metricset['ElementUsageUniqueName']:
                        self.measurs.append(metricset['ElementUsageUniqueName'])

    def readaction(self):
        if 'ClickActions' in abj.keys():
            for action in abj['ClickActions']:
                if 'ActionType' in action.keys():
                    self.action.append(action['ActionType'])
                elif 'Script' in action.keys():
                    self.action.append('Script')

    def readscript(self):
        if 'ClickActions' in abj.keys():
            for action in abj['ClickActions']:
                if 'Script' in action.keys():
                    self.script.append(action['Script'])


    def readname(self):
        if 'Name' in abj.keys():
            self.name.append(abj['Name'])


def Metric():
    a = Chart()
    a.readmetricset()
    return a.metricsets
def Measure():
    b=Chart()
    b.readmeasures()
    return b.measurs
def Action():
    c= Chart()
    c.readaction()
    return c.action

def Name():
    d= Chart()
    d.readname()
    return d.name
def Script():
    e = Chart()
    e.readscript()
    return e.script

def chart_detail_df(chart_detail):
    a = Metric()
    b = Measure()
    c = Action()
    d = Name()
    e = Script()
    if len(d) != 0:
        chart.append(d)
        chart.append(a)
        chart.append(b)
        chart.append(c)
        chart.append(e)
    else:
        pass
    if len(chart) != 0:
        df1 = pd.DataFrame(chart).T
        df1[0].fillna(method='ffill', inplace=True)
        df1[1].fillna(method='ffill', inplace=True)

        chart_detail = pd.concat([df1, chart_detail], ignore_index=True)
    return chart_detail

def image_detail_df(image_detail):

    c = Action()
    d = Name()
    e = Script()
    if len(d) != 0:
        image.append(d)
        image.append(c)
        image.append(e)
    else:
        pass
    if len(image) != 0:
        df1 = pd.DataFrame(image).T
        df1[0].fillna(method='ffill', inplace=True)
        df1[1].fillna(method='ffill', inplace=True)

        image_detail = pd.concat([df1, image_detail], ignore_index=True)
    return image_detail

def label_detail_df(label_detail):

    c = Action()
    d = Name()
    e = Script()
    if len(d) != 0:
        label.append(d)
        label.append(c)
        label.append(e)
    else:
        pass
    if len(label) != 0:
        df1 = pd.DataFrame(label).T
        df1[0].fillna(method='ffill', inplace=True)
        df1[1].fillna(method='ffill', inplace=True)

        label_detail = pd.concat([df1, label_detail], ignore_index=True)
    return label_detail

def insert_to_word(dataframe,table):
    # Table.Cell(1, col + 1).Range.Text = str(df.columns[col])
    for i in range(dataframe.shape[0]):
        for j in range(dataframe.shape[-1]):
            table.cell(0, j ).text = str(dataframe.columns[j])
            if dataframe.values[i,j] != None:
                    table.cell(i+1,j).text = str(dataframe.values[i,j])

if __name__ == "__main__":
    chart_table = pd.DataFrame()
    image_table = pd.DataFrame()
    label_table = pd.DataFrame()



    for abj in adaptor:
        if abj['UIClassName'] == 'dundas.view.controls.Chart':
            chart = []
            chart_table = chart_detail_df(chart_table)
        elif abj['UIClassName'] == 'dundas.view.controls.Image':
            image = []
            image_table = image_detail_df(image_table)
        elif abj['UIClassName'] == 'dundas.view.controls.Label':
            label = []
            label_table = label_detail_df(label_table)
        # elif abj['UIClassName'] == 'dundas.view.controls.Frame':
        #     chart = [['Frame']]
        #     object_table = object_detail_df(object_table)

    index_image = image_table[image_table[1].values ==None].index
    image_table.drop(index_image,inplace=True)
    index_label = label_table[label_table[1].values ==None].index
    label_table.drop(index_label,inplace=True)
    chart_table[5]=None
    image_table[3] = None
    label_table[3] = None
    chart_table.columns = ['Name','MetricSet','Measure','Action','Script','توضیحات']
    image_table.columns = ['Name','Action','Script','توضیحات']
    label_table.columns = ['Name','Action','Script','توضیحات']
    # print(label_table.to_string())
    # print(image_table.to_string())

    chart_table.to_excel("C:/Users/Lenovo/Desktop/measure.xlsx")

    document = docx.Document('C:/Users/Lenovo/Desktop/measure.docx')

    section = document.sections[0]
    section.page_height = Mm(210)
    section.page_width = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(12.7)
    section.footer_distance = Mm(12.7)

    document.add_paragraph('Chart')
    table = document.add_table(chart_table.shape[0]+1, chart_table.shape[1],style=document.styles['table1'])
    document.add_paragraph('image')
    table2 = document.add_table(image_table.shape[0] + 1, image_table.shape[1],style=document.styles['table1'])
    document.add_paragraph('label')
    table3 = document.add_table(label_table.shape[0] + 1, label_table.shape[1],style=document.styles['table1'])
    table.allow_autofit = True
    table.autofit = True
    table2.allow_autofit = True
    table2.autofit = True
    table3.allow_autofit = True
    table3.autofit = True

    insert_to_word(chart_table,table)
    insert_to_word(image_table,table2)
    insert_to_word(label_table,table3)

    # save the doc
    document.save('C:/Users/Lenovo/Desktop/measure1.docx')
